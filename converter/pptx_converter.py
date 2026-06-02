from pptx import Presentation

def convert_pptx(file_path: str, **kwargs) -> str:
    """
    Converts a PowerPoint file to Markdown.
    Reads slides top-to-bottom to preserve natural reading order.
    """
    prs = Presentation(file_path)
    markdown_sections = []

   
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_md = [f"## Slide {slide_num}"]
        
        
        sorted_shapes = sorted(slide.shapes, key=lambda s: getattr(s, 'top', 0) or 0)
        
        for shape in sorted_shapes:
            if shape.has_text_frame:
                text_md = _shape_to_md(shape)
                if text_md:
                    slide_md.append(text_md)
                    
            elif shape.has_table:
                table_md = _table_to_md(shape.table)
                if table_md:
                    slide_md.append(table_md)
                    
        markdown_sections.append("\n\n".join(slide_md))

    return "\n\n---\n\n".join(markdown_sections)


def _shape_to_md(shape) -> str:
    """Helper: converts a text shape into Markdown bullet points."""
    text_chunks = []
    
    for paragraph in shape.text_frame.paragraphs:
        if not paragraph.text.strip():
            continue
            
        indent = "  " * paragraph.level
  
        if getattr(shape, 'shape_type', None) == 14:
            text_chunks.append(f"### {paragraph.text.strip()}")
        else:
            text_chunks.append(f"{indent}- {paragraph.text.strip()}")
            
    return "\n".join(text_chunks)


def _table_to_md(table) -> str:
    """Helper: converts a pptx Table object to a GFM Markdown table."""
    if not table.rows:
        return ""

    table_md = []
    
    for row_index, row in enumerate(table.rows):
        row_data = [cell.text_frame.text.replace('\n', ' ').strip() for cell in row.cells]
        
        row_string = "| " + " | ".join(row_data) + " |"
        table_md.append(row_string)
        
        if row_index == 0:
            separator = "|-" + "-|-".join(["-" * max(3, len(col)) for col in row_data]) + "-|"
            table_md.append(separator)

    return "\n".join(table_md)