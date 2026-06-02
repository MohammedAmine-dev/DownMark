from docx import Document

HEADING_MAP = {
    'Heading 1': '# ',
    'Heading 2': '## ',
    'Heading 3': '### ',
    'Heading 4': '#### ',
    'Heading 5': '##### ',
    'Heading 6': '###### ',
}

def convert_docx(file_path: str, **kwargs) -> str:
    """
    Converts a Word .docx file to Markdown.
    """
    doc = Document(file_path)
    markdown_sections = []

    for para in doc.paragraphs:
        md_text = _paragraph_to_md(para)
        if md_text.strip(): 
            markdown_sections.append(md_text)

    for table in doc.tables:
        md_table = _table_to_md(table)
        markdown_sections.append(md_table)

    return "\n\n".join(markdown_sections)


def _paragraph_to_md(paragraph) -> str:
    """Helper: converts a single paragraph object to a Markdown line."""
    style_name = paragraph.style.name
    
    if 'List' in style_name:
        prefix = '- '
    elif style_name in HEADING_MAP:
        prefix = HEADING_MAP[style_name]
    else:
        prefix = ''

    text_chunks = []
    for run in paragraph.runs:
        text = run.text
        if not text.strip():
            text_chunks.append(text)
            continue
            
        if run.bold:
            text = f"**{text}**"
        if run.italic:
            text = f"*{text}*"
            
        text_chunks.append(text)

    final_text = prefix + "".join(text_chunks)
    return final_text


def _table_to_md(table) -> str:
    """Helper: converts a Table object to a  Markdown table."""
    if not table.rows:
        return ""

    table_md = []
    
    for row_index, row in enumerate(table.rows):
        row_data = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
        
        row_string = "| " + " | ".join(row_data) + " |"
        table_md.append(row_string)
        
        if row_index == 0:
            separator = "|-" + "-|-".join(["-" * len(col) for col in row_data]) + "-|"
            table_md.append(separator)

    return "\n".join(table_md)